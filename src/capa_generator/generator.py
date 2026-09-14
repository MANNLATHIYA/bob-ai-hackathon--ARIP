from __future__ import annotations

from datetime import date, timedelta
from io import BytesIO
from xml.sax.saxutils import escape

from docx import Document
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

from src.api.models import Deviation
from src.deviation_engine.llm import LLMClient, get_llm


TEMPLATES = {
    "MISSED_VISIT": ("Scheduling or participant-contact workflow failure", "Contact participant and assess safety promptly", "Add automated reminders and weekly visit reconciliation", "Site Coordinator"),
    "OUT_OF_WINDOW": ("Visit scheduling did not account for the protocol window", "Document impact and complete outstanding procedures", "Configure protocol-window alerts in the visit calendar", "Site Coordinator"),
    "DOSE_ERROR": ("Dose verification controls were insufficient", "Assess participant and notify investigator/medical monitor", "Introduce independent dose verification and retraining", "Principal Investigator"),
    "WRONG_ROUTE": ("Administration checklist was not followed", "Evaluate participant safety and report per safety plan", "Require route read-back and competency validation", "Principal Investigator"),
    "BANNED_MEDICATION": ("Medication reconciliation did not identify prohibited therapy", "Escalate to medical monitor and evaluate participant safety", "Reconcile medications at every visit with blacklist alerts", "Site Pharmacist"),
    "MISSING_ASSESSMENT": ("Visit checklist was incomplete", "Complete assessment if clinically appropriate and document impact", "Use a locked visit checklist with completion review", "Site Coordinator"),
}


class CAPAGenerator:
    def __init__(self, llm: LLMClient | None = None): self.llm = llm or get_llm()

    async def generate(self, deviation: Deviation) -> str:
        root, corrective, preventive, owner = TEMPLATES.get(deviation.kind, ("Process or documentation control gap", "Review and correct the affected record", "Retrain staff and add quality review", "Site Quality Lead"))
        response = await self.llm.complete_json("Return JSON with root_cause, corrective_action, preventive_action, owner. Keep claims hypothetical and require human approval.", deviation.model_dump_json())
        if response.get("root_cause"):
            root, corrective, preventive, owner = (response.get("root_cause", root), response.get("corrective_action", corrective), response.get("preventive_action", preventive), response.get("owner", owner))
        due = date.today() + timedelta(days=15 if deviation.severity == "MAJOR" else 30)
        return f"""# Corrective and Preventive Action Report

**Report status:** Draft — requires investigator and Quality approval  
**Site:** {deviation.site_id}  
**Deviation:** {deviation.id} ({deviation.kind})  
**Severity:** {deviation.severity}  
**Protocol clause:** {deviation.protocol_clause}  
**GCP basis:** ICH E6(R2) §4.5.1 (protocol compliance), §4.5.3 (document deviations), and §5.20 (noncompliance).  
**Target closure:** {due.isoformat()}

## Deviation summary
{deviation.description}

## Classification rationale
{deviation.rationale}

## Root-cause hypothesis
{root}. Confirm through documented site review before closure.

## Corrective action
{corrective}.

## Preventive action
{preventive}.

## Recommended owner
{owner}

## Effectiveness check
Quality reviewer to confirm action completion and verify no recurrence across the next three applicable visits.
"""


def to_docx(markdown: str) -> bytes:
    doc = Document()
    for line in markdown.splitlines():
        if line.startswith("# "): doc.add_heading(line[2:], level=1)
        elif line.startswith("## "): doc.add_heading(line[3:], level=2)
        elif line.strip(): doc.add_paragraph(line.replace("**", ""))
    stream = BytesIO(); doc.save(stream); return stream.getvalue()


def to_pdf(markdown: str) -> bytes:
    stream = BytesIO(); styles = getSampleStyleSheet(); story = []
    for line in markdown.splitlines():
        if not line.strip(): story.append(Spacer(1, 7)); continue
        style = styles["Title"] if line.startswith("# ") else styles["Heading2"] if line.startswith("## ") else styles["BodyText"]
        story.append(Paragraph(escape(line.lstrip("# ").replace("**", "")), style))
    SimpleDocTemplate(stream, pagesize=LETTER).build(story); return stream.getvalue()

